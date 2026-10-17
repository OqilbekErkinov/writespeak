"""Reports for teachers and learning centers (bot/handlers/group.py and the
teacher web cabinet): a group's Excel workbook, a group summary PDF and a
monthly progress PDF about one student, written for their parents.

Everything here covers only work a student submitted after joining the group
(db/crud.py's group section). Texts are in the teacher's language (uz/ru/en).
"""
from __future__ import annotations

import base64
import io
from datetime import datetime, timedelta

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from weasyprint import HTML  # noqa: E402

from db.crud import criteria_averages, criteria_tip, summarize_works  # noqa: E402
from services.pdf.report_builder import TEMPLATE_DIR, _env  # noqa: E402

_TASHKENT = timedelta(hours=5)
CRITERIA_ORDER = ("TA", "CC", "LR", "GRA", "FC", "Pronunciation")
CRITERIA_SHORT = {"TA": "TA/TR", "CC": "CC", "LR": "LR", "GRA": "GRA", "FC": "FC", "Pronunciation": "P"}
CRITERIA_NAMES = {
    "uz": {
        "TA": "Topshiriqqa javob (TA/TR)", "CC": "Izchillik (CC)", "LR": "Lug'at (LR)",
        "GRA": "Grammatika (GRA)", "FC": "Ravonlik (FC)", "Pronunciation": "Talaffuz",
    },
    "ru": {
        "TA": "Выполнение задания (TA/TR)", "CC": "Связность (CC)", "LR": "Лексика (LR)",
        "GRA": "Грамматика (GRA)", "FC": "Беглость (FC)", "Pronunciation": "Произношение",
    },
    "en": {
        "TA": "Task Achievement / Response (TA/TR)", "CC": "Coherence & Cohesion (CC)",
        "LR": "Lexical Resource (LR)", "GRA": "Grammar (GRA)", "FC": "Fluency & Coherence (FC)",
        "Pronunciation": "Pronunciation",
    },
}

LABELS = {
    "uz": {
        "students": "Talabalar", "works": "Ishlar", "assignments": "Vazifalar",
        "name": "Ism", "joined": "Qo'shilgan", "writing": "Writing", "speaking": "Speaking",
        "count": "soni", "avg": "o'rtacha", "weakest": "Eng zaif mezon", "last": "Oxirgi faollik",
        "this_week": "Bu hafta", "date": "Sana", "student": "Talaba", "type": "Turi", "band": "Ball",
        "assignment": "Vazifa", "paid": "To'lagan", "paid_self": "talaba", "paid_sponsor": "guruh",
        "due": "Muddat", "overall": "Umumiy (W+S)", "teacher_comment": "O'qituvchi izohi",
        "group_title": "Guruh hisoboti", "generated": "Tayyorlangan sana",
        "active_week": "7 kunda faol", "checks": "Tekshiruvlar", "members": "Talabalar",
        "group_weakest": "Guruhning zaif mezoni", "no_work": "hali ish yo'q",
        "parent_title": "Oylik rivojlanish hisoboti", "period": "Davr", "group": "Guruh",
        "teacher": "O'qituvchi", "this_month": "Shu oy", "prev_month": "O'tgan oy",
        "band_trend": "Ballar dinamikasi", "criteria": "Mezonlar bo'yicha o'rtacha ball",
        "criterion": "Mezon", "assignments_done": "Bajarilgan vazifalar", "recommendation": "Tavsiya",
        "comments": "O'qituvchi izohlari", "no_works_month": "Bu oyda tekshiruv topshirilmagan.",
        "disclaimer": "Ballar IELTS Band Descriptors'ning 4 mezoni bo'yicha qo'yilgan. Bandzilla mustaqil xizmat; "
                      "rasmiy IELTS natijasini faqat imtihon markazi beradi.",
        "of": "dan",
    },
    "ru": {
        "students": "Ученики", "works": "Работы", "assignments": "Задания",
        "name": "Имя", "joined": "В группе с", "writing": "Writing", "speaking": "Speaking",
        "count": "кол-во", "avg": "средний", "weakest": "Слабый критерий", "last": "Последняя активность",
        "this_week": "За неделю", "date": "Дата", "student": "Ученик", "type": "Тип", "band": "Балл",
        "assignment": "Задание", "paid": "Оплатил", "paid_self": "ученик", "paid_sponsor": "группа",
        "due": "Срок", "overall": "Итог (W+S)", "teacher_comment": "Комментарий преподавателя",
        "group_title": "Отчёт по группе", "generated": "Дата отчёта",
        "active_week": "Активны за 7 дней", "checks": "Проверок", "members": "Учеников",
        "group_weakest": "Слабый критерий группы", "no_work": "пока нет работ",
        "parent_title": "Ежемесячный отчёт о прогрессе", "period": "Период", "group": "Группа",
        "teacher": "Преподаватель", "this_month": "Этот месяц", "prev_month": "Прошлый месяц",
        "band_trend": "Динамика баллов", "criteria": "Средний балл по критериям",
        "criterion": "Критерий", "assignments_done": "Выполнено заданий", "recommendation": "Рекомендация",
        "comments": "Комментарии преподавателя", "no_works_month": "В этом месяце проверок не было.",
        "disclaimer": "Баллы выставлены по 4 критериям IELTS Band Descriptors. Bandzilla — независимый сервис; "
                      "официальный результат IELTS выдаёт только экзаменационный центр.",
        "of": "из",
    },
    "en": {
        "students": "Students", "works": "Works", "assignments": "Assignments",
        "name": "Name", "joined": "Joined", "writing": "Writing", "speaking": "Speaking",
        "count": "count", "avg": "average", "weakest": "Weakest criterion", "last": "Last active",
        "this_week": "This week", "date": "Date", "student": "Student", "type": "Type", "band": "Band",
        "assignment": "Assignment", "paid": "Paid by", "paid_self": "student", "paid_sponsor": "group",
        "due": "Deadline", "overall": "Overall (W+S)", "teacher_comment": "Teacher's comment",
        "group_title": "Group report", "generated": "Generated on",
        "active_week": "Active in 7 days", "checks": "Checks", "members": "Students",
        "group_weakest": "Group's weakest criterion", "no_work": "no works yet",
        "parent_title": "Monthly progress report", "period": "Period", "group": "Group",
        "teacher": "Teacher", "this_month": "This month", "prev_month": "Last month",
        "band_trend": "Band trend", "criteria": "Average band by criterion",
        "criterion": "Criterion", "assignments_done": "Assignments done", "recommendation": "Recommendation",
        "comments": "Teacher's comments", "no_works_month": "No checks were submitted this month.",
        "disclaimer": "Bands are given on the 4 criteria of the IELTS Band Descriptors. Bandzilla is an "
                      "independent service; only a test centre issues an official IELTS result.",
        "of": "of",
    },
}


def _local(value: datetime | None) -> datetime | None:
    return value + _TASHKENT if value else None


def _day(value: datetime | None) -> str:
    return _local(value).strftime("%d.%m.%Y") if value else "—"


def _band(value: float | None) -> str:
    return "—" if value is None else f"{value:.1f}"


# --- Excel ---------------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="062018")  # Bandzilla Deep Ink


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "B2"
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1])) for r in rows if index - 1 < len(r) and r[index - 1] is not None])
        ws.column_dimensions[get_column_letter(index)].width = min(max(10, width + 2), 48)
    return ws


def build_group_excel(group_name: str, members: list[tuple], assignments: list[dict], lang: str = "uz") -> bytes:
    """`members`: (user, joined_at, works) from db.crud.get_group_works;
    `assignments`: dicts with "assignment" and "results" (db.crud.get_assignment_results)."""
    L = LABELS.get(lang, LABELS["uz"])
    week_start = datetime.now().astimezone() - timedelta(days=7)
    part_titles = {}
    assignment_of_part = {}
    for item in assignments:
        for part in item["results"]["parts"]:
            assignment_of_part[part.id] = item["assignment"].title
            part_titles[part.id] = part.module.replace("_", " ").title()

    student_rows, work_rows = [], []
    for user, joined_at, works in members:
        s = summarize_works(works)
        criteria = s["criteria"]
        student_rows.append(
            [user.full_name, _day(joined_at), s["writing"], s["writing_avg"], s["speaking"], s["speaking_avg"]]
            + [round(criteria[k], 2) if k in criteria else None for k in CRITERIA_ORDER]
            + [CRITERIA_NAMES[lang if lang in CRITERIA_NAMES else "uz"].get(s["weakest"][0], s["weakest"][0]) if s["weakest"] else None,
               _day(s["last_active"]), sum(1 for w in works if w["created_at"] >= week_start)]
        )
        for work in works:
            work_rows.append(
                [_local(work["created_at"]).strftime("%d.%m.%Y %H:%M"), user.full_name, work["label"], work["band"]]
                + [work["criteria"].get(k) for k in CRITERIA_ORDER]
                + [assignment_of_part.get(work["assignment_part_id"], ""),
                   L["paid_sponsor"] if work["paid_by_id"] else L["paid_self"]]
            )
    work_rows.sort(key=lambda r: r[0], reverse=True)

    wb = Workbook()
    _sheet(
        wb, L["students"],
        [L["name"], L["joined"], f"{L['writing']} ({L['count']})", f"{L['writing']} ({L['avg']})",
         f"{L['speaking']} ({L['count']})", f"{L['speaking']} ({L['avg']})"]
        + [CRITERIA_SHORT[k] for k in CRITERIA_ORDER] + [L["weakest"], L["last"], L["this_week"]],
        student_rows, first=True,
    )
    _sheet(
        wb, L["works"],
        [L["date"], L["student"], L["type"], L["band"]] + [CRITERIA_SHORT[k] for k in CRITERIA_ORDER]
        + [L["assignment"], L["paid"]],
        work_rows,
    )
    assignment_rows = []
    for item in assignments:
        assignment, results = item["assignment"], item["results"]
        for row in results["rows"]:
            bands = [b[0] if b else None for b in row["bands"]]
            line = [assignment.title, _day(assignment.due_at), row["user"].full_name] + bands
            if assignment.kind == "mock":
                line.append(row["mock"]["overall"] if row["mock"] else None)
            assignment_rows.append(line)
    width = max((len(item["results"]["parts"]) for item in assignments), default=1)
    _sheet(
        wb, L["assignments"],
        [L["assignment"], L["due"], L["student"]] + [f"{L['band']} {i + 1}" for i in range(width)] + [L["overall"]],
        assignment_rows,
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDFs ---------------------------------------------------------------------------------------


def band_chart(works: list[dict], title: str, daily_average: bool = False) -> str | None:
    """A band-over-time line per skill, as a data URI - None with fewer than two
    points to draw. `daily_average`: one point per day (a whole group's work)."""
    series = []
    for kind, label, color in (("w", "Writing", "#0A7C4A"), ("s", "Speaking", "#FF7A93")):  # Moss, Tongue Pink
        points = sorted((_local(w["created_at"]).replace(tzinfo=None), w["band"]) for w in works if w["kind"] == kind)
        if daily_average:
            days: dict = {}
            for when, band in points:
                days.setdefault(when.date(), []).append(band)
            points = [(datetime(d.year, d.month, d.day), sum(b) / len(b)) for d, b in sorted(days.items())]
        if points:
            series.append((label, color, points))
    times = [when for _, _, points in series for when, _ in points]
    if len(times) < 2:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 2.6), dpi=150)
    for label, color, points in series:
        ax.plot([p[0] for p in points], [p[1] for p in points], marker="o", linewidth=2, label=label, color=color)
    ax.set_ylim(3.5, 9)
    ax.set_title(title, fontsize=10, loc="left")
    ax.grid(alpha=0.25)
    ax.legend(fontsize=8, loc="lower right")
    ax.tick_params(labelsize=8)
    short = max(times) - min(times) < timedelta(days=2)  # everything within a day or two: show the hour too
    ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=3, maxticks=5 if short else 8))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d.%m %H:%M" if short else "%d.%m"))
    fig.tight_layout()
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png")
    plt.close(fig)
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


def _render(template: str, **context) -> bytes:
    html = _env.get_template(template).render(**context)
    return HTML(string=html, base_url=str(TEMPLATE_DIR)).write_pdf()


def build_group_pdf(
    group_name: str, teacher_name: str, overview: dict, members: list[tuple], assignments: list[dict], lang: str = "uz"
) -> bytes:
    L = LABELS.get(lang, LABELS["uz"])
    names = CRITERIA_NAMES.get(lang, CRITERIA_NAMES["uz"])
    rows = []
    for user, joined_at, works in members:
        s = summarize_works(works)
        rows.append({
            "name": user.full_name, "joined": _day(joined_at), "writing": s["writing"],
            "writing_avg": _band(s["writing_avg"]), "speaking": s["speaking"], "speaking_avg": _band(s["speaking_avg"]),
            "weakest": names.get(s["weakest"][0], s["weakest"][0]) if s["weakest"] else "—",
            "last": _day(s["last_active"]), "total": s["total"],
        })
    rows.sort(key=lambda r: -r["total"])
    summary = overview["summary"]
    all_works = [w for _, _, works in members for w in works]
    assignment_rows = [
        {"title": item["assignment"].title, "due": _day(item["assignment"].due_at),
         "complete": item["results"]["complete"], "members": len(item["results"]["rows"])}
        for item in assignments
    ]
    return _render(
        "group_report.html",
        L=L, lang=lang, group_name=group_name, teacher_name=teacher_name, generated=_day(datetime.now().astimezone()),
        overview=overview, rows=rows, assignments=assignment_rows,
        writing_avg=_band(summary["writing_avg"]), speaking_avg=_band(summary["speaking_avg"]),
        weakest=names.get(summary["weakest"][0], summary["weakest"][0]) if summary["weakest"] else "—",
        chart=band_chart(all_works, L["band_trend"], daily_average=True),
    )


def build_parent_report_pdf(
    *,
    student_name: str,
    group_name: str,
    teacher_name: str,
    works_month: list[dict],
    works_prev: list[dict],
    comments: list[tuple[dict, list]],
    assignments_done: int,
    assignments_total: int,
    period_start: datetime,
    period_end: datetime,
    lang: str = "uz",
) -> bytes:
    """`comments`: (work, [SubmissionComment]) pairs from this month."""
    L = LABELS.get(lang, LABELS["uz"])
    names = CRITERIA_NAMES.get(lang, CRITERIA_NAMES["uz"])
    month, prev = summarize_works(works_month), summarize_works(works_prev)
    now_c, prev_c = criteria_averages(works_month), criteria_averages(works_prev)
    criteria_rows = [
        {"name": names.get(k, k), "now": _band(now_c.get(k)), "prev": _band(prev_c.get(k)),
         "delta": (now_c[k] - prev_c[k]) if k in now_c and k in prev_c else None}
        for k in CRITERIA_ORDER if k in now_c or k in prev_c
    ]
    tip = None
    if month["weakest"]:
        key = month["weakest"][0]
        tip = criteria_tip(key, lang) or None
    comment_rows = [
        {"date": _day(work["created_at"]), "label": work["label"], "band": _band(work["band"]),
         "text": c.text, "teacher_band": _band(c.band) if c.band is not None else None}
        for work, items in comments for c in items
    ]
    return _render(
        "parent_report.html",
        L=L, lang=lang, student_name=student_name, group_name=group_name, teacher_name=teacher_name,
        period=f"{_day(period_start)} – {_day(period_end)}",
        month=month, prev=prev, band=_band, criteria_rows=criteria_rows,
        weakest=names.get(month["weakest"][0], month["weakest"][0]) if month["weakest"] else None,
        tip=tip, comments=comment_rows, assignments_done=assignments_done, assignments_total=assignments_total,
        chart=band_chart(works_month, L["band_trend"]),
    )
