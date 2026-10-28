"""The admin panel's control room (the user's request, 2026-10-18: "we must
control the bot fully from the panel"):

- /settings      the test window on/off and its testers (bot/utils/runtime.py);
- /broadcast     a message to all users or a segment, in their language;
- /support       the students' messages to the admins and the answers;
- /errors        the bot's, the importer's and the panel's errors;
- /export/*.xlsx users, payments and checks for Excel;
- block / unblock a user;
- /admins        separate logins (owner only), /audit what was done by whom.

Wired into webapp/main.py like webapp/cabinet.py (`templates` is set there).
Every change made here is written to the audit log (admin_actions).
"""
from __future__ import annotations

import asyncio
import hashlib
import hmac
import html
import io
import json
import logging
import secrets
import time
from datetime import datetime, timedelta, timezone

from aiogram import Bot
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramForbiddenError, TelegramRetryAfter
from fastapi import APIRouter, Form, Request
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from sqlalchemy import and_, func, select, update

from bot.config import settings
from bot.utils import checks_back, crm, runtime
from db.database import get_session
from db.models import (
    AdminAction,
    ChartKey,
    PaymentKind,
    PaymentRequest,
    PaymentStatus,
    SupportMessage,
    User,
    UserStatus,
    WebAdmin,
)
from webapp import insights

logger = logging.getLogger(__name__)
router = APIRouter()
templates = None  # set by webapp/main.py

TESTERS_LIST_KEY = "runtime:testers"  # the testers list, kept while the window is off
BROADCAST_KEY = "broadcast:{id}"
BROADCAST_IDS = "broadcast:ids"
BROADCAST_PAUSE = 0.05  # 20 messages a second at most - under Telegram's limit


# --- who is signed in, and the audit log ------------------------------------------------------------


def logged_in(request: Request) -> bool:
    return bool(request.session.get("logged_in"))


def admin_name(request: Request) -> str:
    return request.session.get("admin") or settings.web_admin_username


def is_owner(request: Request) -> bool:
    # sessions from before separate logins were the owner's (the .env login)
    return bool(request.session.get("owner", "admin" not in request.session))


def guard(request: Request, owner: bool = False):
    """A redirect when the request may not see the page, else None."""
    if not logged_in(request):
        return RedirectResponse("/login", status_code=303)
    if owner and not is_owner(request):
        return RedirectResponse("/?denied=1", status_code=303)
    return None


async def audit(request: Request | str, action: str, target_user_id: int | None = None, details: str | None = None) -> None:
    who = request if isinstance(request, str) else admin_name(request)
    try:
        async with get_session() as session:
            session.add(AdminAction(admin=who, action=action, target_user_id=target_user_id, details=details))
            await session.commit()
    except Exception as e:  # noqa: BLE001 - the action itself has happened
        logger.warning("Audit entry %s by %s not saved: %r", action, who, e)


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode(), salt=salt, n=2 ** 14, r=8, p=1)
    return f"scrypt${salt.hex()}${digest.hex()}"


def check_password(password: str, stored: str) -> bool:
    try:
        scheme, salt, digest = stored.split("$")
        if scheme != "scrypt":
            return False
        got = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=2 ** 14, r=8, p=1)
        return hmac.compare_digest(got.hex(), digest)
    except Exception:  # noqa: BLE001
        return False


def _same(a: str, b: str) -> bool:
    return hmac.compare_digest(a.encode(), b.encode())


def _pw_tag(password_hash: str) -> str:
    """Kept in a web admin's session: a new password (or a disabled login)
    ends the sessions signed in before it. Not the hash itself - the
    session cookie is signed, not encrypted."""
    return hashlib.sha256(password_hash.encode()).hexdigest()[:16]


async def authenticate(username: str, password: str) -> dict | None:
    """The session's sign-in keys for a valid login: the .env owner (always
    works - the break-glass login) or an active web admin."""
    if _same(username, settings.web_admin_username) and _same(password, settings.web_admin_password):
        return {"admin": settings.web_admin_username, "owner": True}
    async with get_session() as session:
        row = (await session.execute(select(WebAdmin).where(WebAdmin.username == username.lower()))).scalar_one_or_none()
        if row is None or not row.active or not check_password(password, row.password_hash):
            return None
        row.last_login_at = datetime.now(timezone.utc)
        await session.commit()
        return {"admin": row.username, "owner": False, "pw": _pw_tag(row.password_hash)}


_valid_cache: dict[str, tuple[float, str | None]] = {}  # username -> (checked at, its pw tag; None = disabled)


async def session_valid(request: Request) -> bool:
    """False for a web admin who was disabled or given a new password since
    signing in (checked at most every 15 seconds)."""
    if is_owner(request):
        return True
    username = request.session.get("admin") or ""
    cached = _valid_cache.get(username)
    if cached is None or time.monotonic() - cached[0] > 15:
        try:
            async with get_session() as session:
                row = (await session.execute(select(WebAdmin).where(WebAdmin.username == username))).scalar_one_or_none()
        except Exception:  # noqa: BLE001 - the database is down: the pages fail anyway
            return True
        cached = (time.monotonic(), _pw_tag(row.password_hash) if row is not None and row.active else None)
        _valid_cache[username] = cached
    return cached[1] is not None and cached[1] == request.session.get("pw")


def new_bot() -> Bot:
    return Bot(settings.telegram_bot_token, default=DefaultBotProperties(parse_mode=ParseMode.HTML))


# Pages that never show the top bar - no badges to count for them.
NAV_SKIP = ("/login", "/logout", "/cabinet", "/reports/", "/export/", "/support/file/", "/broadcast/", "/payments/",
            "/questions/", "/chart-image/")


async def nav_counts(request: Request) -> dict:
    """The top bar's badges, and who is signed in."""
    out = {"support": 0, "payments": 0, "charts": 0, "errors": await errors_today(), "window": False,
           "owner": is_owner(request), "admin": admin_name(request)}
    try:
        async with get_session() as session:
            out["support"] = (await session.execute(select(func.count()).select_from(SupportMessage).where(
                SupportMessage.direction == "in", SupportMessage.read.is_(False)))).scalar() or 0
            out["payments"] = (await session.execute(select(func.count()).select_from(PaymentRequest).where(
                PaymentRequest.status == PaymentStatus.pending))).scalar() or 0
            out["charts"] = (await session.execute(select(func.count()).select_from(ChartKey).where(
                ChartKey.status == "auto"))).scalar() or 0
        out["window"] = await runtime.window_on()
    except Exception as e:  # noqa: BLE001 - the page itself still opens
        logger.info("Nav badges not counted: %r", e)
    return out


_jobs: set[asyncio.Task] = set()


def _later(coro) -> None:
    """Runs `coro` after the response, in this process; its failure is logged."""
    async def run():
        try:
            await coro
        except Exception:
            logger.exception("Panel background job failed")
    task = asyncio.get_running_loop().create_task(run())
    _jobs.add(task)  # a task nothing refers to may be garbage-collected mid-way
    task.add_done_callback(_jobs.discard)


# --- 1. the test window ------------------------------------------------------------------------------------


async def _testers_list() -> set[int]:
    raw = await crm.redis().get(TESTERS_LIST_KEY)
    if raw is None:  # first use: the .env list
        return set(settings.checks_only_for_ids)
    return runtime._parse(raw)


async def _save_testers_list(ids: set[int]) -> None:
    await crm.redis().set(TESTERS_LIST_KEY, ",".join(str(i) for i in sorted(ids)))


async def _notify_back_job() -> None:
    bot = new_bot()
    try:
        sent = await checks_back.notify_back(bot)
        logger.info("Test window off from the panel: %d users told checks are back", sent)
    finally:
        await bot.session.close()


@router.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request, done: str | None = None):
    if (r := guard(request, owner=True)) is not None:
        return r
    on = await runtime.window_on()
    listed = await _testers_list()
    async with get_session() as session:
        users = {u.id: u for u in (await session.execute(select(User).where(User.id.in_(listed)))).scalars()} if listed else {}
    flash = {
        "on": "Test rejimi yoqildi: endi faqat sinovchilar va adminlar tekshira oladi.",
        "off": "Tekshiruvlar hamma uchun ochildi. Kutayotganlarga «Bandzilla yana ishlayapti» xabari yuborilmoqda.",
        "added": "Sinovchi qo'shildi.", "removed": "Sinovchi olib tashlandi.", "bad": "Bunday foydalanuvchi topilmadi.",
        "empty": "Avval kamida bitta sinovchi qo'shing.",
    }.get(done or "")
    return templates.TemplateResponse(request, "settings.html", {
        "active": "settings", "on": on, "testers": sorted(listed), "users": users, "waiting": await checks_back.waiting_count(),
        "flash": flash, "flash_bad": done in ("bad", "empty"), "env_on": bool(settings.checks_only_for_ids),
    })


@router.post("/settings/test-window")
async def settings_test_window(request: Request, action: str = Form(...), user_id: str = Form("")):
    if (r := guard(request, owner=True)) is not None:
        return r
    listed = await _testers_list()
    on = await runtime.window_on()
    done = action
    if action in ("add", "remove"):
        raw = user_id.strip().lstrip("@")
        async with get_session() as session:
            if raw.isdigit():
                uid = int(raw)
                exists = await session.get(User, uid) is not None
            else:  # an @username
                uid = (await session.execute(select(User.id).where(func.lower(User.username) == raw.lower())
                                             .limit(1))).scalar_one_or_none()
                exists = uid is not None
        if action == "add" and not exists:
            return RedirectResponse("/settings?done=bad", status_code=303)
        listed = listed | {uid} if action == "add" else listed - {uid}
        await _save_testers_list(listed)
        if on:
            await runtime.set_testers(listed)
        done = "added" if action == "add" else "removed"
        await audit(request, f"tester_{action}", uid)
    elif action == "on":
        if not listed:
            return RedirectResponse("/settings?done=empty", status_code=303)
        await runtime.set_testers(listed)
        await audit(request, "test_window_on", details=f"{len(listed)} sinovchi")
    elif action == "off":
        await _save_testers_list(listed)
        await runtime.set_testers(set())
        await audit(request, "test_window_off", details=f"{await checks_back.waiting_count()} kishi kutayotgan edi")
        _later(_notify_back_job())
    return RedirectResponse(f"/settings?done={done}", status_code=303)


# --- 2. broadcast --------------------------------------------------------------------------------------------

AUDIENCES = {
    "all": "Barcha foydalanuvchilar", "active": "7 kunda faol bo'lganlar", "checked": "Tekshiruv qilganlar",
    "nochecks": "Hali tekshiruv qilmaganlar", "paying": "To'lov qilganlar", "new": "Bugun qo'shilganlar",
    "teachers": "O'qituvchilar", "admins": "Faqat adminlar (sinov)",
}
LANG_FILTERS = {"any": "Hamma tillar", "uz": "🇺🇿 O'zbekcha", "ru": "🇷🇺 Ruscha", "en": "🇬🇧 Inglizcha"}


async def _audience(session, audience: str, lang: str) -> list[User]:
    if audience == "admins":
        rows = (await session.execute(select(User).where(User.id.in_(settings.admin_ids)))).scalars().all()
        return list(rows)
    data = await insights.users_overview(session, None, audience if audience in insights.SEGMENTS else "all", limit=10 ** 6)
    users = [r["user"] for r in data["rows"]]
    users = [u for u in users if u.status != UserStatus.denied and u.id not in settings.admin_ids]
    if lang != "any":
        users = [u for u in users if (u.language or "uz") == lang]
    return users


async def _broadcast_job(bid: str, users: list[tuple[int, str]], texts: dict[str, str]) -> None:
    key = BROADCAST_KEY.format(id=bid)
    r = crm.redis()
    bot = new_bot()
    sent = failed = blocked = 0
    try:
        for i, (uid, lang) in enumerate(users, 1):
            text = texts.get(lang) or texts["uz"]
            try:
                try:
                    await bot.send_message(uid, text, disable_web_page_preview=True)
                except TelegramRetryAfter as e:
                    await asyncio.sleep(e.retry_after)
                    await bot.send_message(uid, text, disable_web_page_preview=True)
                sent += 1
            except TelegramForbiddenError:
                blocked += 1
            except Exception as e:  # noqa: BLE001
                failed += 1
                logger.info("Broadcast %s to %s failed: %r", bid, uid, e)
            if i % 10 == 0 or i == len(users):
                await r.hset(key, mapping={"sent": sent, "failed": failed, "blocked": blocked, "done": i,
                                         "updated": int(time.time())})
            await asyncio.sleep(BROADCAST_PAUSE)
        await r.hset(key, mapping={"status": "done", "finished": int(time.time())})
    except Exception:
        await r.hset(key, mapping={"status": "stopped"})
        raise
    finally:
        await bot.session.close()
    await crm.post("support", f"📣 <b>Ommaviy xabar</b> #{bid} tugadi: {sent} ta yetkazildi, {blocked} ta botni bloklagan, "
                              f"{failed} ta xato.")


def _checked(data: dict) -> dict:
    if data.get("status") == "running" and time.time() - int(data.get("updated") or data.get("started") or 0) > 180:
        data["status"] = "stopped"  # the panel restarted mid-way (a deploy): the job is gone
    return data


async def _broadcasts() -> list[dict]:
    r = crm.redis()
    out = []
    for bid in await r.lrange(BROADCAST_IDS, 0, 19):
        data = await r.hgetall(BROADCAST_KEY.format(id=bid))
        if data:
            data["id"] = bid
            data["when"] = insights.fmt(datetime.fromtimestamp(int(data.get("started", 0)), timezone.utc))
            out.append(_checked(data))
    return out


@router.get("/broadcast", response_class=HTMLResponse)
async def broadcast_page(request: Request, audience: str = "all", lang: str = "any", started: str | None = None):
    if (r := guard(request)) is not None:
        return r
    audience = audience if audience in AUDIENCES else "all"
    lang = lang if lang in LANG_FILTERS else "any"
    async with get_session() as session:
        users = await _audience(session, audience, lang)
    langs = {code: sum(1 for u in users if (u.language or "uz") == code) for code in ("uz", "ru", "en")}
    return templates.TemplateResponse(request, "broadcast.html", {
        "active": "broadcast", "audiences": AUDIENCES, "lang_filters": LANG_FILTERS, "audience": audience, "lang": lang,
        "count": len(users), "langs": langs, "history": await _broadcasts(), "started": started,
    })


@router.post("/broadcast/send")
async def broadcast_send(request: Request, audience: str = Form(...), lang: str = Form("any"),
                         text_uz: str = Form(""), text_ru: str = Form(""), text_en: str = Form("")):
    if (r := guard(request)) is not None:
        return r
    texts = {k: html.escape(v.strip()[:3500], quote=False) for k, v in (("uz", text_uz), ("ru", text_ru), ("en", text_en))}
    if not texts["uz"]:
        return RedirectResponse(f"/broadcast?audience={audience}&lang={lang}&started=empty", status_code=303)
    async with get_session() as session:
        users = await _audience(session, audience if audience in AUDIENCES else "all", lang if lang in LANG_FILTERS else "any")
    targets = [(u.id, u.language or "uz") for u in users]
    bid = datetime.now(timezone.utc).strftime("%y%m%d%H%M%S") + secrets.token_hex(2)
    r = crm.redis()
    await r.hset(BROADCAST_KEY.format(id=bid), mapping={
        "status": "running", "total": len(targets), "sent": 0, "failed": 0, "blocked": 0, "done": 0,
        "audience": AUDIENCES.get(audience, audience), "lang": LANG_FILTERS.get(lang, lang), "admin": admin_name(request),
        "preview": text_uz.strip()[:160], "started": int(time.time()),
    })
    await r.lpush(BROADCAST_IDS, bid)
    await r.ltrim(BROADCAST_IDS, 0, 49)
    await audit(request, "broadcast", details=f"#{bid} · {AUDIENCES.get(audience)} · {LANG_FILTERS.get(lang)} · "
                                              f"{len(targets)} kishi · {text_uz.strip()[:200]}")
    _later(_broadcast_job(bid, targets, texts))
    return RedirectResponse(f"/broadcast?started={bid}", status_code=303)


@router.get("/broadcast/{bid}.json")
async def broadcast_status(request: Request, bid: str):
    if not logged_in(request):
        return JSONResponse({}, status_code=401)
    return JSONResponse(_checked(await crm.redis().hgetall(BROADCAST_KEY.format(id=bid))))


# --- 3. support inbox ----------------------------------------------------------------------------------------


@router.get("/support", response_class=HTMLResponse)
async def support_page(request: Request, f: str = "all"):
    if (r := guard(request)) is not None:
        return r
    async with get_session() as session:
        last = select(SupportMessage.user_id, func.max(SupportMessage.id).label("mid"),
                      func.count().filter(and_(SupportMessage.direction == "in", SupportMessage.read.is_(False)))
                      .label("unread"), func.count().label("n")).group_by(SupportMessage.user_id).subquery()
        rows = (await session.execute(
            select(SupportMessage, User, last.c.unread, last.c.n)
            .join(last, last.c.mid == SupportMessage.id).join(User, User.id == SupportMessage.user_id)
            .order_by(SupportMessage.created_at.desc()))).all()
    threads = [{"m": m, "user": u, "unread": unread, "n": n, "ago": insights.ago(m.created_at)} for m, u, unread, n in rows]
    if f == "unread":
        threads = [x for x in threads if x["unread"]]
    elif f == "waiting":
        threads = [x for x in threads if x["m"].direction == "in"]
    return templates.TemplateResponse(request, "support.html", {"active": "support", "threads": threads, "f": f})


@router.get("/support/{user_id}", response_class=HTMLResponse)
async def support_thread(request: Request, user_id: int, sent: str | None = None):
    if (r := guard(request)) is not None:
        return r
    async with get_session() as session:
        user = await session.get(User, user_id)
        if user is None:
            return RedirectResponse("/support", status_code=303)
        messages = (await session.execute(select(SupportMessage).where(SupportMessage.user_id == user_id)
                                          .order_by(SupportMessage.created_at))).scalars().all()
        await session.execute(update(SupportMessage).where(SupportMessage.user_id == user_id,
                                                           SupportMessage.direction == "in").values(read=True))
        await session.commit()
    for m in messages:
        m.when = insights.fmt(m.created_at)
    flash = {"1": "Javob yuborildi.", "0": "Javob yetkazilmadi: foydalanuvchi botni bloklagan bo'lishi mumkin.",
             "empty": "Javob bo'sh."}.get(sent or "")
    return templates.TemplateResponse(request, "support_thread.html", {
        "active": "support", "user": user, "messages": messages, "flash": flash, "flash_bad": sent in ("0", "empty"),
    })


@router.get("/support/file/{message_id}")
async def support_file(request: Request, message_id: int):
    """A photo from the conversation, fetched from Telegram with the bot's token."""
    if not logged_in(request):
        return Response(status_code=401)
    async with get_session() as session:
        m = await session.get(SupportMessage, message_id)
    if m is None or not m.file_id:
        return Response(status_code=404)
    bot = Bot(settings.telegram_bot_token)
    try:
        file = await bot.get_file(m.file_id)
        data = (await bot.download_file(file.file_path)).read()
    except Exception:  # noqa: BLE001
        return Response(status_code=404)
    finally:
        await bot.session.close()
    return Response(data, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=86400"})


# --- 4. errors -----------------------------------------------------------------------------------------------


async def errors_today() -> int:
    """Errors since midnight in Tashkent (a line repeated within 10 minutes is one)."""
    since = insights.day_start_utc().timestamp()
    try:
        raw = await crm.redis().lrange(crm.ERRORS_LOG_KEY, 0, crm.ERRORS_LOG_CAP - 1)
    except Exception:  # noqa: BLE001
        return 0
    n = 0
    for item in raw:
        try:
            e = json.loads(item)
        except ValueError:
            continue
        if e.get("at", 0) < since:
            break  # newest first
        n += e.get("level") == "ERROR"
    return n


@router.get("/errors", response_class=HTMLResponse)
async def errors_page(request: Request, src: str = "all"):
    if (r := guard(request)) is not None:
        return r
    raw = await crm.redis().lrange(crm.ERRORS_LOG_KEY, 0, crm.ERRORS_LOG_CAP - 1)
    entries = []
    for item in raw:
        try:
            e = json.loads(item)
        except ValueError:
            continue
        e["when"] = insights.fmt(datetime.fromtimestamp(e.get("at", 0), timezone.utc))
        e["ago"] = insights.ago(datetime.fromtimestamp(e.get("at", 0), timezone.utc))
        entries.append(e)
    sources = sorted({e.get("process", "bot") for e in entries})
    if src != "all":
        entries = [e for e in entries if e.get("process", "bot") == src]
    day = datetime.now(timezone.utc) - timedelta(days=1)
    return templates.TemplateResponse(request, "errors.html", {
        "active": "errors", "entries": entries, "src": src, "sources": sources, "owner": is_owner(request),
        "last24": sum(1 for e in entries if e.get("at", 0) >= day.timestamp()),
    })


@router.post("/errors/clear")
async def errors_clear(request: Request):
    if (r := guard(request, owner=True)) is not None:
        return r
    await crm.redis().delete(crm.ERRORS_LOG_KEY)
    await audit(request, "errors_cleared")
    return RedirectResponse("/errors", status_code=303)


# --- 5. block / unblock --------------------------------------------------------------------------------------


@router.post("/users/{user_id}/block")
async def user_block(request: Request, user_id: int, action: str = Form(...)):
    if (r := guard(request)) is not None:
        return r
    if user_id in settings.admin_ids:
        return RedirectResponse(f"/users/{user_id}?blocked=admin", status_code=303)
    async with get_session() as session:
        user = await session.get(User, user_id)
        if user is None:
            return RedirectResponse("/users", status_code=303)
        user.status = UserStatus.denied if action == "block" else UserStatus.approved
        name = user.full_name
        await session.commit()
    await audit(request, "user_blocked" if action == "block" else "user_unblocked", user_id)
    await crm.post("users", f"{'⛔ <b>Bloklandi</b>' if action == 'block' else '✅ <b>Blokdan chiqarildi</b>'}: "
                            f"{html.escape(name, quote=False)} (<code>{user_id}</code>) · web panel, {html.escape(admin_name(request))}")
    return RedirectResponse(f"/users/{user_id}?blocked={action}", status_code=303)


# --- 6. Excel ------------------------------------------------------------------------------------------------


def _cell(value):
    """A value Excel won't run: names and texts come from Telegram users, so a
    leading = + - @ (a formula) is quoted."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _xlsx(title: str, headers: list[str], rows: list[list]) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="062018")
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([_cell(v) for v in row])
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:500]])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"Bandzilla_{title}_{datetime.now(timezone.utc) + timedelta(hours=5):%Y-%m-%d}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _local(dt):
    return (dt + timedelta(hours=5)).replace(tzinfo=None) if dt else None


@router.get("/export/users.xlsx")
async def export_users(request: Request):
    if (r := guard(request)) is not None:
        return r
    async with get_session() as session:
        data = await insights.users_overview(session, None, "all", limit=10 ** 6)
    rows = [[r["user"].id, r["user"].full_name, r["user"].username, r["user"].language, r["user"].role.value,
             r["user"].status.value, r["user"].coin_balance, r["writing"], r["speaking"], r["paid"],
             _local(r["user"].created_at), _local(r["last_active"]), r["user"].referred_by] for r in data["rows"]]
    await audit(request, "export", details="users.xlsx")
    return _xlsx("Foydalanuvchilar", ["ID", "Ism", "Username", "Til", "Rol", "Holat", "Coin", "Writing", "Speaking",
                                      "To'lagan (so'm)", "Qo'shilgan", "Oxirgi faollik", "Taklif qilgan ID"], rows)


@router.get("/export/payments.xlsx")
async def export_payments(request: Request):
    if (r := guard(request)) is not None:
        return r
    async with get_session() as session:
        payments = (await session.execute(select(PaymentRequest).order_by(PaymentRequest.created_at.desc()))).scalars().all()
        ids = {p.user_id for p in payments}
        users = {u.id: u for u in (await session.execute(select(User).where(User.id.in_(ids)))).scalars()} if ids else {}
    kinds = {PaymentKind.credits: "coin", PaymentKind.gift: "sovg'a", PaymentKind.subscription: "obuna"}
    rows = [[p.id, _local(p.created_at), p.user_id, users[p.user_id].full_name if p.user_id in users else "",
             (users[p.user_id].username or "") if p.user_id in users else "", kinds.get(p.kind, p.kind.value), p.coins,
             p.amount, p.status.value, _local(p.reviewed_at), "web panel" if p.reviewed_by == -1 else p.reviewed_by]
            for p in payments]
    await audit(request, "export", details="payments.xlsx")
    return _xlsx("Tolovlar", ["#", "Sana", "Foydalanuvchi ID", "Ism", "Username", "Turi", "Coin", "Summa (so'm)",
                              "Holat", "Ko'rib chiqildi", "Kim"], rows)


@router.get("/export/checks.xlsx")
async def export_checks(request: Request):
    if (r := guard(request)) is not None:
        return r
    async with get_session() as session:
        data = await insights.all_checks(session, "all", "all", limit=100000)
    rows = [[c["kind"].upper(), c["id"], _local(c["at"]), c["user_id"], c["user"].full_name if c["user"] else "",
             c["label"], c["band"], c["criteria"], c["coins"], c["paid"], round(c["cost"], 4)] for c in data["rows"]]
    await audit(request, "export", details="checks.xlsx")
    return _xlsx("Tekshiruvlar", ["Modul", "#", "Sana", "Foydalanuvchi ID", "Ism", "Turi", "Band", "Mezonlar", "Coin",
                                  "To'lov", "AI xarajat ($)"], rows)


# --- 7. admins and the audit log -----------------------------------------------------------------------------


@router.get("/admins", response_class=HTMLResponse)
async def admins_page(request: Request, done: str | None = None):
    if (r := guard(request, owner=True)) is not None:
        return r
    async with get_session() as session:
        admins = (await session.execute(select(WebAdmin).order_by(WebAdmin.created_at))).scalars().all()
    password = request.session.pop("new_password", None)
    flash = {"created": "Admin qo'shildi.", "reset": "Yangi parol berildi.", "toggled": "Holat o'zgartirildi.",
             "exists": "Bu login band.", "bad": "Login: 3-30 ta lotin harfi, raqam yoki _ ."}.get(done or "")
    for a in admins:
        a.last = insights.ago(a.last_login_at) if a.last_login_at else "hali kirmagan"
    return templates.TemplateResponse(request, "admins.html", {
        "active": "admins", "admins": admins, "flash": flash, "flash_bad": done in ("exists", "bad"), "password": password,
        "owner_login": settings.web_admin_username,
    })


@router.post("/admins/create")
async def admins_create(request: Request, username: str = Form(...), name: str = Form(...)):
    if (r := guard(request, owner=True)) is not None:
        return r
    username = username.strip().lower()
    if not (3 <= len(username) <= 30 and all(c.isascii() and (c.isalnum() or c == "_") for c in username)) \
            or username == settings.web_admin_username:
        return RedirectResponse("/admins?done=bad", status_code=303)
    password = secrets.token_urlsafe(10)
    async with get_session() as session:
        if (await session.execute(select(WebAdmin).where(WebAdmin.username == username))).scalar_one_or_none():
            return RedirectResponse("/admins?done=exists", status_code=303)
        session.add(WebAdmin(username=username, name=name.strip()[:100] or username, password_hash=hash_password(password)))
        await session.commit()
    request.session["new_password"] = {"username": username, "password": password}
    await audit(request, "admin_created", details=username)
    return RedirectResponse("/admins?done=created", status_code=303)


@router.post("/admins/{admin_id}/{action}")
async def admins_change(request: Request, admin_id: int, action: str):
    if (r := guard(request, owner=True)) is not None:
        return r
    async with get_session() as session:
        a = await session.get(WebAdmin, admin_id)
        if a is None or action not in ("reset", "toggle"):
            return RedirectResponse("/admins", status_code=303)
        if action == "reset":
            password = secrets.token_urlsafe(10)
            a.password_hash = hash_password(password)
            request.session["new_password"] = {"username": a.username, "password": password}
        else:
            a.active = not a.active
        username, active = a.username, a.active
        await session.commit()
    _valid_cache.pop(username, None)
    await audit(request, "admin_password_reset" if action == "reset" else ("admin_enabled" if active else "admin_disabled"),
                details=username)
    return RedirectResponse(f"/admins?done={'reset' if action == 'reset' else 'toggled'}", status_code=303)


ACTION_NAMES = {
    "login": "Kirdi", "login_failed": "Kirish xato", "payment_approve": "To'lovni tasdiqladi",
    "payment_reject": "To'lovni rad etdi", "coins": "Coin o'zgartirdi", "message": "Xabar yozdi",
    "broadcast": "Ommaviy xabar", "user_blocked": "Blokladi", "user_unblocked": "Blokdan chiqardi",
    "test_window_on": "Test rejimini yoqdi", "test_window_off": "Test rejimini o'chirdi", "tester_add": "Sinovchi qo'shdi",
    "tester_remove": "Sinovchini olib tashladi", "admin_created": "Admin qo'shdi", "admin_password_reset": "Parolni yangiladi",
    "admin_enabled": "Adminni yoqdi", "admin_disabled": "Adminni o'chirdi", "export": "Excel yuklab oldi",
    "errors_cleared": "Xatolarni tozaladi", "center_created": "Markaz yaratdi", "center_deleted": "Markazni o'chirdi",
    "teacher_revoked": "O'qituvchilikni oldi",
}


@router.get("/audit", response_class=HTMLResponse)
async def audit_page(request: Request, admin: str = "", action: str = ""):
    if (r := guard(request, owner=True)) is not None:
        return r
    async with get_session() as session:
        stmt = select(AdminAction).order_by(AdminAction.created_at.desc()).limit(400)
        if admin:
            stmt = stmt.where(AdminAction.admin == admin)
        if action:
            stmt = stmt.where(AdminAction.action == action)
        rows = (await session.execute(stmt)).scalars().all()
        admins = (await session.execute(select(AdminAction.admin).distinct())).scalars().all()
        ids = {a.target_user_id for a in rows if a.target_user_id}
        users = {u.id: u for u in (await session.execute(select(User).where(User.id.in_(ids)))).scalars()} if ids else {}
    for a in rows:
        a.when = insights.fmt(a.created_at)
        a.label = ACTION_NAMES.get(a.action, a.action)
    return templates.TemplateResponse(request, "audit.html", {
        "active": "audit", "rows": rows, "admins": sorted(admins), "admin": admin, "action": action,
        "actions": ACTION_NAMES, "users": users,
    })
