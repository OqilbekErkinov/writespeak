"""Task 1 chart pictures, from Telegram to the vision model and back
(2026-10-19): image checks, the chart's data sheet, and the bot's Writing
flows with a chart.

Offline - OpenAI and Telegram are fakes, so nothing is sent and nothing is
paid for. Part 1 (pictures, errors, the data sheet) runs anywhere the bot's
requirements are installed. Part 2 drives the real bot handlers, so it needs a
THROWAWAY database and Redis, never production's - inside the bot image:

    DATABASE_URL=... REDIS_URL=... OPENAI_API_KEY=sk-dummy python tests/task1_images_test.py

Part 2 flushes that Redis and writes Writing submissions into that database.
"""
import asyncio
import io
import logging
import os
import re
import sys
from datetime import datetime, timezone
from types import SimpleNamespace as NS

import httpx
import openai
from PIL import Image, ImageDraw

fails: list[str] = []
RUN = os.urandom(3).hex()  # every run's answers are new essays, never "the same essay again"


def check(name: str, ok, detail="") -> None:
    print(("ok   " if ok else "FAIL ") + name + (f"  [{detail}]" if not ok and detail != "" else ""))
    if not ok:
        fails.append(name)


# --- pictures -----------------------------------------------------------------------------------------

BARS = [("Sports", 27, 12), ("Computer games", 18, 15), ("Music", 20, 20), ("Shopping", 35, 7)]


def chart(fmt: str = "PNG", size=(1100, 700), mode: str = "RGB", **save) -> bytes:
    """A bar chart like report #62's: boys and girls, four activities."""
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    w, h = size
    base = h - max(20, h // 9)
    scale = max((base - 20) / 40, 0.5)
    d.line((70, base, w - 40, base), fill="black", width=2)
    d.line((70, 40, 70, base), fill="black", width=2)
    for v in range(0, 40, 5):
        d.text((25, base - v * scale - 6), f"{v}%", fill="black")
    step = (w - 140) / len(BARS)
    for i, (label, boys, girls) in enumerate(BARS):
        x = 100 + i * step
        d.rectangle((x, base - boys * scale, x + step * 0.3, base), fill=(10, 124, 74))
        d.rectangle((x + step * 0.35, base - girls * scale, x + step * 0.65, base), fill=(240, 190, 0))
        d.text((x, base + 12), label, fill="black")
    if mode == "RGBA":
        img = img.convert("RGBA")
        img.putalpha(Image.new("L", size, 0).point(lambda _: 0))  # a transparent background...
        ImageDraw.Draw(img).rectangle((100, 100, 300, 300), fill=(0, 0, 0, 255))  # ...with something on it
    elif mode != "RGB":
        img = img.convert(mode)
    buf = io.BytesIO()
    img.save(buf, format=fmt, **save)
    return buf.getvalue()


def part1_pictures() -> None:
    from services.storage.file_storage import (
        MIN_IMAGE_SIDE, VISION_MAX_BYTES, VISION_MAX_SIDE, BadImage, image_mime, prepare_vision_image)

    def opened(data: bytes) -> Image.Image:
        return Image.open(io.BytesIO(data))

    for fmt, mime, ext in (("PNG", "image/png", ".png"), ("JPEG", "image/jpeg", ".jpg"), ("WEBP", "image/webp", ".webp")):
        raw = chart(fmt)
        check(f"a {fmt} chart goes to the model byte for byte", prepare_vision_image(raw) == (raw, mime, ext))
        check(f"its type is told by its bytes: {mime}", image_mime(raw, "x.jpg") == mime)

    data, mime, ext = prepare_vision_image(chart("PNG", mode="RGBA"))
    img = opened(data)
    check("a transparent PNG is flattened onto white (not black)", mime == "image/png" and img.mode == "RGB"
          and img.getpixel((5, 5)) == (255, 255, 255) and img.getpixel((150, 150)) == (0, 0, 0))

    data, mime, _ = prepare_vision_image(chart("PNG", size=(4000, 2500)))
    check("a huge chart is shrunk to the model's 2048 px, still lossless PNG",
          mime == "image/png" and max(opened(data).size) == VISION_MAX_SIDE)

    noisy = Image.frombytes("RGB", (2400, 2400), os.urandom(2400 * 2400 * 3))
    buf = io.BytesIO()
    noisy.save(buf, format="PNG")
    data, mime, _ = prepare_vision_image(buf.getvalue())
    check("a photo-like PNG too big even at 2048 px becomes a JPEG under the size limit",
          mime == "image/jpeg" and len(data) <= VISION_MAX_BYTES, (mime, len(data)))

    data, mime, _ = prepare_vision_image(chart("JPEG", mode="CMYK"))
    check("a CMYK JPEG becomes an RGB JPEG", mime == "image/jpeg" and opened(data).mode == "RGB")

    exif = Image.Exif()
    exif[0x0112] = 6  # "rotate 90" - how phone cameras save a portrait shot
    data, mime, _ = prepare_vision_image(chart("JPEG", exif=exif))
    check("a phone photo is turned upright (EXIF)", opened(data).size == (700, 1100), opened(data).size)

    frames = [Image.open(io.BytesIO(chart("PNG"))).convert("P"), Image.new("P", (1100, 700))]
    buf = io.BytesIO()
    frames[0].save(buf, format="GIF", save_all=True, append_images=frames[1:])
    data, mime, _ = prepare_vision_image(buf.getvalue())
    check("an animated GIF becomes a still PNG (the model takes no animation)", mime == "image/png")

    data, mime, _ = prepare_vision_image(chart("BMP"))
    check("a BMP becomes a PNG", mime == "image/png")

    for name, raw, reason in (
        ("not a picture at all", b"%PDF-1.4 definitely not an image", "not_image"),
        ("a truncated PNG", chart("PNG")[:4000], "not_image"),
        ("a HEIC file (no decoder)", b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00mif1heic" + b"\x00" * 200, "not_image"),
        ("a thumbnail", chart("PNG", size=(MIN_IMAGE_SIDE - 50, 150)), "too_small"),
    ):
        try:
            prepare_vision_image(raw)
            check(f"{name} is refused", False, "accepted")
        except BadImage as e:
            check(f"{name} is refused ({reason})", e.reason == reason, e.reason)

    check("an unknown file falls back to its name's type", image_mime(b"????", "a.webp") == "image/webp")


def bad_request(code: str, message: str) -> openai.BadRequestError:
    request = httpx.Request("POST", "https://api.openai.com/v1/chat/completions")
    return openai.BadRequestError(message, response=httpx.Response(400, request=request),
                                  body={"code": code, "message": message, "type": "invalid_request_error"})


def server_error() -> openai.InternalServerError:
    request = httpx.Request("POST", "https://api.openai.com/v1/chat/completions")
    return openai.InternalServerError("upstream failed", response=httpx.Response(500, request=request), body=None)


def reading(kind: str = "clear"):
    from services.ai.schemas import VisualFact, VisualReading

    facts = []
    for label, boys, girls in BARS:
        facts.append(VisualFact(label=f"{label} - Boys", value=f"about {boys}% ({boys - 2}-{boys + 2}%)", status="estimated"))
        facts.append(VisualFact(label=f"{label} - Girls", value=f"about {girls}% ({girls - 2}-{girls + 2}%)", status="estimated"))
    note = ""
    if kind == "partly_unclear":
        facts[-1] = VisualFact(label="Shopping - Girls", value="the top of the bar is cropped", status="unclear")
        note = "the girls' shopping bar is cropped"
    if kind == "unreadable":
        facts, note = [], "the picture shows only the question text, no chart"
    return VisualReading(
        visual_type="bar chart", subject="favourite leisure activities of teenagers in Canada, by gender",
        units="%", facts=facts,
        main_features=["Shopping is the boys' top activity and the girls' least popular.",
                       "Boys lead in sports and computer games."],
        cautions=["Music: boys and girls look equal, about 20%."],
        readability=kind, readability_note=note,
    )


async def part1_reading() -> None:
    from services.ai import visual_reading
    from services.ai.openai_client import is_image_error

    check("OpenAI's 'unsupported image' is an image error",
          is_image_error(bad_request("invalid_image_format", "You uploaded an unsupported image.")))
    check("'image_parse_error' is an image error", is_image_error(bad_request("image_parse_error", "Bad image")))
    check("another 400 is not", not is_image_error(bad_request("context_length_exceeded", "Too long")))
    check("a server error is not", not is_image_error(server_error()))

    sheet = visual_reading.sheet_text(reading("clear"))
    check("the sheet lists estimated values with their range", "Sports - Boys: about 27% (25-29%) (estimated from the scale)"
          in sheet and "Music: boys and girls look equal" in sheet)
    check("an unreadable reading gives no sheet", visual_reading.sheet_text(reading("unreadable")) == "")
    check("a partly unclear one says what can't be read", "(can't be read)" in visual_reading.sheet_text(reading("partly_unclear")))

    real = visual_reading.parse_flex
    images = [(chart("PNG"), "image/png")]
    outcomes = {}

    async def fake(**kw):
        outcome = outcomes["next"]
        if isinstance(outcome, BaseException):
            raise outcome
        outcomes["messages"] = kw["messages"]
        return NS(choices=[NS(message=NS(parsed=outcome, refusal="I can't help" if outcome is None else None))],
                  usage=None, model="gpt-5")

    visual_reading.parse_flex = fake
    try:
        outcomes["next"] = reading("clear")
        got = await visual_reading.read_visual("The chart shows ...", images)
        image_parts = [p for p in outcomes["messages"][1]["content"] if p["type"] == "image_url"]
        check("the chart reaches the vision model as a PNG data URL in high detail",
              got is not None and len(image_parts) == 1 and image_parts[0]["image_url"]["detail"] == "high"
              and image_parts[0]["image_url"]["url"].startswith("data:image/png;base64,"))
        outcomes["next"] = bad_request("invalid_image_format", "You uploaded an unsupported image.")
        try:
            await visual_reading.read_visual("q", images)
            check("a picture OpenAI can't open stops the check", False, "no exception")
        except visual_reading.ChartImageRejected:
            check("a picture OpenAI can't open stops the check (asked again)", True)
        outcomes["next"] = server_error()
        check("a failed reading call: the check goes on without a sheet", await visual_reading.read_visual("q", images) is None)
        outcomes["next"] = None
        check("an empty / refused reading: no sheet", await visual_reading.read_visual("q", images) is None)
        outcomes["next"] = openai.LengthFinishReasonError(completion=NS(usage=None, model="gpt-5"))
        check("a reading cut off at its token limit: no sheet", await visual_reading.read_visual("q", images) is None)
        empty = reading("clear")
        empty.facts = []
        outcomes["next"] = empty
        check("a 'clear' reading with no data is not trusted", await visual_reading.read_visual("q", images) is None)
    finally:
        visual_reading.parse_flex = real


# --- part 2: the bot, with fake OpenAI and Telegram ----------------------------------------------------

ANSWER = (
    "The bar chart illustrates the favourite leisure activities of teenage boys and girls in Canada.\n\n"
    "Overall, boys is more interested in shopping and sports, whereas girls show a striking contrast in their choices.\n\n"
    "Around 27% of boys prefer sports, compared with 12% of girls. Shopping is the most popular pastime for boys "
    "at 35%, while very few girls go shopping, at 7%.\n\n"
    "Computer games attract 18% of boys and 15% of girls, and music is equally popular, at 20% for both groups."
)


def sample_text(words: int) -> list[str]:
    bank = ("the graph compares how young people in one nation spend their free time, split by gender. "
            "in general, the most noticeable difference appears in one pursuit, where one group dominates. "
            "by contrast, another activity attracts nearly identical shares from each group of respondents. "
            "the remaining figures sit between these extremes and change only modestly across the categories. ").split()
    out, i = [], 0
    while len(out) < words:
        out.append(bank[i % len(bank)])
        i += 1
    text = " ".join(out)
    quarter = len(text) // 4
    paragraphs = [text[:quarter], text[quarter:2 * quarter], text[2 * quarter:3 * quarter], text[3 * quarter:]]
    return [p.strip().capitalize() + "." for p in paragraphs]


class FakeAI:
    """Answers every OpenAI call the Writing check makes, and keeps them."""

    def __init__(self):
        self.calls: list[tuple[str, list]] = []
        self.readings: list = []  # what the next chart readings return (or raise)
        self.answer = ANSWER
        self.fail_scoring = False
        self.feedback_none = False

    def names(self) -> list[str]:
        return [n for n, _ in self.calls]

    async def parse(self, **kw):
        from services.ai import schemas

        fmt = kw["response_format"]
        name = fmt.__name__
        self.calls.append((name, kw["messages"]))
        if name == "VisualReading":
            outcome = self.readings.pop(0) if self.readings else reading("clear")
            if isinstance(outcome, BaseException):
                raise outcome
            parsed = outcome
        elif name == "CriterionScores":
            if self.fail_scoring:
                raise server_error()
            parsed = schemas.CriterionScores(TR_or_TA=6, CC=6, LR=6, GRA=6)
        elif name == "WritingFeedback":
            parsed = None if self.feedback_none else self.feedback()
        elif name == "MissedMistakes":
            parsed = schemas.MissedMistakes(mistakes=[])
        elif name == "Verdicts":
            text = _text_of(kw["messages"])
            listed = re.findall(r"^(\d+)\. «(.*?)»", text, flags=re.M)
            # A careful checker, as the real one is asked to be: a figure inside
            # the data sheet's range is not a mistake.
            parsed = schemas.Verdicts(verdicts=[
                schemas.Verdict(id=int(i), keep=not ("DATA SHEET" in text and "7%" in wrong), reason="checked")
                for i, wrong in listed])
        elif name == "GeneratedSample":
            task1 = "Academic Writing Task 1" in kw["messages"][0]["content"]
            parsed = schemas.GeneratedSample(paragraphs=sample_text(185 if task1 else 320), shows=[])
        elif name == "_QualityCheck":
            parsed = fmt(ok=True, problem="")
        else:
            raise AssertionError(f"unexpected OpenAI call {name}")
        return NS(choices=[NS(message=NS(parsed=parsed, refusal=None, content=None))],
                  usage=None, model="gpt-5", service_tier="flex")

    async def create(self, **kw):  # OCR
        self.calls.append(("ocr", kw["messages"]))
        return NS(choices=[NS(message=NS(content="The chart below shows the favourite leisure activities of "
                                                 "teenagers in Canada. Summarise the information."))],
                  usage=None, model="gpt-5")

    def feedback(self):
        from services.ai import schemas

        wrongs = {"at 7%": 1, "boys is": 2}
        essay = []
        for paragraph in self.answer.split("\n\n"):
            segments, rest = [], paragraph
            for wrong, mid in wrongs.items():
                if wrong in rest:
                    before, after = rest.split(wrong, 1)
                    segments += [schemas.EssaySegment(type="text", text=before, fix=None, mistake_id=None),
                                 schemas.EssaySegment(type="error", text=wrong, fix=None, mistake_id=mid)]
                    rest = after
            segments.append(schemas.EssaySegment(type="text", text=rest, fix=None, mistake_id=None))
            essay.append([s for s in segments if s.text])
        return schemas.WritingFeedback(
            mistakes=[
                schemas.Mistake(id=1, wrong="at 7%", correct="at about 8%", why="Girls' shopping is roughly 8%.",
                                criterion="TR", unclear=False, if_you_meant=None),
                schemas.Mistake(id=2, wrong="boys is", correct="boys are", why="A plural subject takes a plural verb.",
                                criterion="GRA", unclear=False, if_you_meant=None),
            ],
            essay=essay,
            criterion_notes=schemas.CriterionNotes(TR_or_TA="Clear overview; figures match the chart.",
                                                   CC="Logical order with clear paragraphs.",
                                                   LR="Accurate comparison vocabulary.", GRA="Mostly accurate grammar."),
            strength=schemas.Strength(text='You compare the groups well: "a striking contrast".',
                                      quoted_phrases=["a striking contrast"]),
            priorities=[schemas.Priority(title=f"Plural verbs {i}", rule="A plural subject takes a plural verb.",
                                         example_wrong="boys is", example_right="boys are", criteria=["GRA"],
                                         mistake_ids=[2]) for i in range(1, 4)],
            vocabulary=[schemas.VocabEntry(term=f"term {i}", meaning="meaning", example="An example sentence.")
                        for i in range(10)],
            checklist=[schemas.ChecklistItem(text="My verbs agree with plural subjects.", ref="Priority 01")
                       for _ in range(5)],
        )


def _text_of(messages: list[dict]) -> str:
    out = []
    for m in messages:
        content = m["content"]
        if isinstance(content, str):
            out.append(content)
        else:
            out += [p["text"] for p in content if p["type"] == "text"]
    return "\n".join(out)


def _images_of(messages: list[dict]) -> list[str]:
    return [p["image_url"]["url"] for m in messages if not isinstance(m["content"], str)
            for p in m["content"] if p["type"] == "image_url"]


async def part2_bot() -> None:
    from aiogram import Bot
    from aiogram.client.session.base import BaseSession
    from aiogram.types import CallbackQuery, Chat, Document, File, Message, PhotoSize, Update, User
    from sqlalchemy import select

    from bot.config import settings
    from bot.i18n import t
    from bot.keyboards.confirm_kb import CB_CONFIRM
    from bot.keyboards.writing_task_kb import CB_NO_CHART, CB_TASK1, CB_TASK2
    from bot.main import build_dispatcher
    from bot.states.writing_states import WritingStates
    from bot.utils.crm import redis
    from db.database import get_session
    from db.models import WritingSubmission
    from services.ai.openai_client import client
    from services.ai.visual_reading import ANSWER_RULES, JUDGE_RULES

    await redis().flushdb()
    fake = FakeAI()
    client.with_options = lambda **kw: client  # the Flex call goes through the same fake
    client.beta.chat.completions.parse = fake.parse
    client.chat.completions.create = fake.create

    files: dict[str, bytes] = {}
    now = datetime.now(timezone.utc)

    class TgSession(BaseSession):
        def __init__(self):
            super().__init__()
            self.calls, self.n = [], 5000

        async def close(self):
            pass

        async def stream_content(self, url, headers=None, timeout=30, chunk_size=65536, raise_for_status=True):
            yield files[url.rsplit("/", 1)[-1]]

        async def make_request(self, bot, method, timeout=None):
            self.calls.append(method)
            self.n += 1
            name = type(method).__name__
            if name == "GetFile":
                return File(file_id=method.file_id, file_unique_id="u" + method.file_id,
                            file_size=len(files[method.file_id]), file_path=f"documents/{method.file_id}")
            if name in ("AnswerCallbackQuery", "DeleteMessage", "EditMessageReplyMarkup", "SetMessageReaction",
                        "SendChatAction"):
                return True
            chat_id = getattr(method, "chat_id", None) or uid
            return Message(message_id=getattr(method, "message_id", None) or self.n, date=now,
                           chat=Chat(id=chat_id, type="private"), text=getattr(method, "text", None) or "x").as_(bot)

    uid = next(iter(settings.admin_ids))  # an admin: never stopped by a test window or a paywall
    session = TgSession()
    bot = Bot("8888548766:TEST", session=session)
    dp = build_dispatcher()
    state = dp.fsm.get_context(bot=bot, chat_id=uid, user_id=uid)
    counter = iter(range(10_000, 100_000))
    sender = User(id=uid, is_bot=False, first_name="T")

    async def send(**kw):
        start = len(session.calls)
        message = Message(message_id=next(counter), date=now, chat=Chat(id=uid, type="private"), from_user=sender, **kw)
        await dp.feed_update(bot, Update(update_id=next(counter), message=message))
        return session.calls[start:]

    async def press(data: str):
        start = len(session.calls)
        query = CallbackQuery(id=str(next(counter)), from_user=sender, chat_instance="ci", data=data,
                              message=Message(message_id=next(counter), date=now, chat=Chat(id=uid, type="private"), text="x"))
        await dp.feed_update(bot, Update(update_id=next(counter), callback_query=query))
        return session.calls[start:]

    def texts(calls) -> list[str]:
        return [c.text for c in calls if type(c).__name__ in ("SendMessage", "EditMessageText")]

    def document_sent(calls):
        return [c for c in calls if type(c).__name__ == "SendDocument"]

    async def lang() -> str:
        from db import crud
        async with get_session() as s:
            return (await crud.get_user(s, uid)).language

    def doc(file_id: str, data: bytes, name: str, mime: str) -> Document:
        files[file_id] = data
        return Document(file_id=file_id, file_unique_id="u" + file_id, file_name=name, mime_type=mime, file_size=len(data))

    async def last_submission():
        async with get_session() as s:
            return (await s.execute(select(WritingSubmission).where(WritingSubmission.user_id == uid)
                                    .order_by(WritingSubmission.id.desc()))).scalars().first()

    async def task1_until_answer(tag: str):
        await press(CB_TASK1)
        await send(text=f"The chart below shows the favourite leisure activities of teenagers in Canada ({tag}). "
                        "Summarise the information by selecting and reporting the main features.")
        check(f"[{tag}] a typed Task 1 question asks for the chart", await state.get_state() == WritingStates.awaiting_chart.state)

    L = await lang()

    # 1. A PNG chart sent as a file: read once, its figures used everywhere.
    await task1_until_answer("png")
    await send(document=doc("chart_png", chart("PNG"), "chart.png", "image/png"))
    check("[png] the chart is taken, the answer asked for", await state.get_state() == WritingStates.awaiting_answer.state)
    fake.calls.clear()
    fake.answer = ANSWER + f" (png {RUN})"
    calls = await send(text=fake.answer)
    check("[png] the report PDF is sent", len(document_sent(calls)) == 1, texts(calls)[-2:])
    names = fake.names()
    check("[png] the chart is read exactly once", names.count("VisualReading") == 1, names)
    check("[png] ...before anything else", names[0] == "VisualReading", names[:2])
    for name in ("CriterionScores", "WritingFeedback", "MissedMistakes"):
        msgs = [m for n, m in fake.calls if n == name]
        check(f"[png] every {name} call judges by the same DATA SHEET and rules",
              msgs and all("DATA SHEET" in _text_of(m) and JUDGE_RULES[:40] in _text_of(m) for m in msgs))
        check(f"[png] ...with the chart attached as PNG", msgs and all(
            _images_of(m) and all(u.startswith("data:image/png;base64,") for u in _images_of(m)) for m in msgs))
    verify = [_text_of(m) for n, m in fake.calls if n == "Verdicts"]
    check("[png] the data mistake ('at 7%') goes to verification, with the sheet",
          any("«at 7%»" in v and "DATA SHEET" in v for v in verify), verify[:1])
    samples = [_text_of(m) for n, m in fake.calls if n in ("GeneratedSample", "_QualityCheck")]
    check("[png] the model answer is written and checked with the same sheet",
          samples and all("DATA SHEET" in s for s in samples) and any(ANSWER_RULES[:40] in s for s in samples))
    sub = await last_submission()
    check("[png] saved: the data sheet and the chart file", sub.answer_text == fake.answer
          and sub.report_data.get("visual", {}).get("visual_type") == "bar chart"
          and sub.original_file_paths.get("chart") and sub.original_file_paths["chart"][0].endswith(".png"))
    marked = [m["wrong"] for m in sub.annotations]
    check("[png] the student's correct '7%' is no longer marked wrong; the real grammar slip is",
          "at 7%" not in marked and "boys is" in marked, marked)
    check("[png] PDF fonts no longer flood the log (the log handler stops DEBUG records)",
          logging.getLogger().handlers and all(h.level >= logging.INFO for h in logging.getLogger().handlers))

    # 1b. The question and its chart in one photo (a screenshot): read (OCR), confirmed, the photo kept as the chart.
    await press(CB_TASK1)
    files["question_photo"] = chart("JPEG", quality=90)
    fake.calls.clear()
    calls = await send(photo=[PhotoSize(file_id="question_photo", file_unique_id="uq", width=1100, height=700)])
    ocr = [m for n, m in fake.calls if n == "ocr"]
    check("[photo question] the photo is read (OCR) as a JPEG and shown to confirm",
          ocr and _images_of(ocr[0])[0].startswith("data:image/jpeg;base64,")
          and await state.get_state() == WritingStates.confirming_prompt.state)
    calls = await press(CB_CONFIRM)
    check("[photo question] confirmed: no separate chart asked - the photo is the chart",
          await state.get_state() == WritingStates.awaiting_answer.state
          and len((await state.get_data()).get("prompt_image_paths") or []) == 1)
    fake.calls.clear()
    fake.answer = ANSWER + f" (photo question {RUN})"
    calls = await send(text=fake.answer)
    check("[photo question] analysed with the chart's data sheet", len(document_sent(calls)) == 1
          and fake.names().count("VisualReading") == 1)

    # 2. A JPEG chart as a Telegram photo; 3. a WebP file.
    for tag, message_kw, mime in (
        ("jpeg", lambda: {"photo": [PhotoSize(file_id="chart_jpg", file_unique_id="uj", width=1100, height=700)]}, "image/jpeg"),
        ("webp", lambda: {"document": doc("chart_webp", chart("WEBP"), "chart.webp", "image/webp")}, "image/webp"),
    ):
        files["chart_jpg"] = chart("JPEG", quality=90)
        await task1_until_answer(tag)
        await send(**message_kw())
        fake.calls.clear()
        fake.answer = ANSWER + f" ({tag} {RUN})"
        calls = await send(text=fake.answer)
        urls = [u for _, m in fake.calls for u in _images_of(m)]
        check(f"[{tag}] analysed: report sent, the chart sent as {mime}", len(document_sent(calls)) == 1
              and urls and all(u.startswith(f"data:{mime};base64,") for u in urls))

    # 4. Pictures that can't be used, at the chart step.
    await task1_until_answer("bad")
    for name, kw, key in (
        ("a damaged PNG", {"document": doc("broken", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64, "chart.png", "image/png")}, "input.bad_image"),
        ("a PDF named .png", {"document": doc("liar", b"%PDF-1.4 1 0 obj", "chart.png", "image/png")}, "input.bad_image"),
        ("an iPhone HEIC file", {"document": doc("heic", b"\x00\x00\x00\x18ftypheic" + b"\x00" * 300, "IMG_1.HEIC", "image/heic")}, "input.bad_image"),
        ("a thumbnail", {"document": doc("tiny", chart("PNG", size=(160, 120)), "tiny.png", "image/png")}, "input.image_too_small"),
        ("a zip file", {"document": doc("zip", b"PK\x03\x04", "files.zip", "application/zip")}, "writing.chart_photo_only"),
        ("text instead of a picture", {"text": "here is the chart"}, "writing.chart_photo_only"),
    ):
        calls = await send(**kw)
        check(f"[bad] {name}: a clear message, still waiting for the chart",
              t(key, L) in texts(calls) and await state.get_state() == WritingStates.awaiting_chart.state, texts(calls))

    # 5. A picture with no usable chart: asked again, the answer kept, then checked.
    await task1_until_answer("unreadable")
    await send(document=doc("text_only", chart("PNG"), "question.png", "image/png"))
    fake.readings = [reading("unreadable"), reading("clear")]
    fake.answer = ANSWER + f" (unreadable {RUN})"
    calls = await send(text=fake.answer)
    data = await state.get_data()
    check("[unreadable] the student is asked for the chart again - nothing charged, no report",
          t("writing.chart_resend", L) in texts(calls) and not document_sent(calls)
          and await state.get_state() == WritingStates.awaiting_chart.state and data.get("rechart"))
    check("[unreadable] ...and the answer is kept", data.get("answer_text") == fake.answer)
    calls = await send(document=doc("chart_png2", chart("PNG"), "chart.png", "image/png"))
    check("[unreadable] the new chart starts the check at once - no second answer asked",
          len(document_sent(calls)) == 1 and t("writing.send_answer", L, min_words=150) not in texts(calls))

    # 6. OpenAI can't open the picture: asked again; going on without it works.
    await task1_until_answer("rejected")
    await send(document=doc("chart_png3", chart("PNG"), "chart.png", "image/png"))
    fake.readings = [bad_request("invalid_image_format", "You uploaded an unsupported image.")]
    fake.calls.clear()
    fake.answer = ANSWER + f" (rejected {RUN})"
    calls = await send(text=fake.answer)
    check("[rejected] OpenAI's 'unsupported image' asks for the chart again, before any scoring",
          t("writing.chart_resend", L) in texts(calls) and "CriterionScores" not in fake.names())
    fake.calls.clear()
    calls = await press(CB_NO_CHART)
    feedback = [_text_of(m) for n, m in fake.calls if n == "WritingFeedback"]
    check("[rejected] 'continue without the chart' checks it as text only", len(document_sent(calls)) == 1
          and "VisualReading" not in fake.names() and feedback and "NO VISUAL WAS PROVIDED" in feedback[0])

    # 7. Partly unclear chart: checked, and the student is told.
    await task1_until_answer("partly")
    await send(document=doc("chart_png4", chart("PNG"), "chart.png", "image/png"))
    fake.readings = [reading("partly_unclear")]
    fake.answer = ANSWER + f" (partly {RUN})"
    calls = await send(text=fake.answer)
    sent = document_sent(calls)
    check("[partly] report sent, and its caption says some values weren't checked",
          sent and t("writing.chart_partly_unclear", L) in (sent[0].caption or ""))

    # 8. OpenAI fails: every scoring run errors / the feedback comes back empty.
    for tag, setup in (("api-error", "fail_scoring"), ("empty", "feedback_none")):
        await task1_until_answer(tag)
        await send(document=doc(f"chart_{tag}", chart("PNG"), "chart.png", "image/png"))
        setattr(fake, setup, True)
        fake.answer = ANSWER + f" ({tag} {RUN})"
        before = (await last_submission()).id
        calls = await send(text=fake.answer)
        setattr(fake, setup, False)
        check(f"[{tag}] the student is told the check failed (no coins), nothing saved",
              t("writing.grading_failed", L) in texts(calls) and not document_sent(calls)
              and (await last_submission()).id == before, texts(calls))

    # 9. Task 2 typed: untouched - no chart reading, no pictures.
    await press(CB_TASK2)
    await send(text="Some people think that advertising does more harm than good. To what extent do you agree?")
    fake.calls.clear()
    fake.answer = ("Advertising is everywhere today. " * 20).strip() + "\n\n" + ("It informs and persuades buyers. " * 20).strip() + f" {RUN}"
    calls = await send(text=fake.answer)
    check("[task2 text] report sent, no chart reading, no pictures", len(document_sent(calls)) == 1
          and "VisualReading" not in fake.names() and not any(_images_of(m) for _, m in fake.calls))

    # 10. Task 2 question as a PDF, answer as a DOCX: still read and checked.
    from docx import Document as Docx
    from weasyprint import HTML

    pdf = HTML(string="<p>Some people believe that cities should ban cars from their centres. Do you agree?</p>").write_pdf()
    word = Docx()
    word.add_paragraph(("Cars make city centres noisy and dangerous. " * 20).strip())
    word.add_paragraph(("Banning them would make streets calmer for everyone. " * 15).strip() + f" {RUN}")
    buf = io.BytesIO()
    word.save(buf)
    await press(CB_TASK2)
    calls = await send(document=doc("q_pdf", pdf, "question.pdf", "application/pdf"))
    check("[pdf] the question is read from the PDF and shown to confirm",
          any("ban cars" in x for x in texts(calls)) and await state.get_state() == WritingStates.confirming_prompt.state)
    await press(CB_CONFIRM)
    calls = await send(document=doc("a_docx", buf.getvalue(), "answer.docx",
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    check("[docx] the answer is read from the Word file", any("Banning them" in x for x in texts(calls)))
    fake.answer = "\n\n".join(p.text for p in Docx(io.BytesIO(buf.getvalue())).paragraphs)
    calls = await press(CB_CONFIRM)
    check("[pdf+docx] report sent", len(document_sent(calls)) == 1)

    await redis().aclose()


async def main() -> int:
    part1_pictures()
    await part1_reading()
    if "--pictures-only" not in sys.argv:
        await part2_bot()
    print("FAILURES:", fails or "none")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
